import json, urllib.request, mimetypes, base64, re
from pathlib import Path
from handlers.utils import log, send_json, safe_path, HOME

def handle_complex_post(h, p, body):
    if p.startswith("/tg_"):
        return handle_telegram(h, p, body)

    if p == "/export":
        return handle_export(h, body)

    if p == "/excel_eval":
        return handle_excel_eval(h, body)

    if p == "/docx_insert_table":
        return handle_docx_table(h, body)

    if p == "/base64":
        return handle_base64_op(h, body)

    return False

def handle_telegram(h, p, body):
    token = body.get("token", "")
    if p == "/tg_get_me":
        try:
            url = f"https://api.telegram.org/bot{token}/getMe"
            req = urllib.request.Request(url, method="GET")
            with urllib.request.urlopen(req, timeout=10) as r:
                resp = json.loads(r.read())
            if resp.get("ok"):
                bot = resp["result"]
                return send_json(h, {"ok": True, "name": bot.get("first_name", ""), "username": bot.get("username", "")})
            return send_json(h, {"error": "Token noto'g'ri"}, 400)
        except Exception as e: return send_json(h, {"error": str(e)}, 500)

    if p == "/tg_send":
        try:
            chat_id = body.get("chat_id", ""); text = body.get("text", "")
            url = f"https://api.telegram.org/bot{token}/sendMessage"
            data = json.dumps({"chat_id": chat_id, "text": text, "parse_mode": "HTML"}).encode()
            req = urllib.request.Request(url, data=data, headers={"Content-Type": "application/json"}, method="POST")
            with urllib.request.urlopen(req, timeout=15) as r:
                resp = json.loads(r.read())
            return send_json(h, {"ok": resp.get("ok", False), "msg_id": resp.get("result", {}).get("message_id")})
        except Exception as e: return send_json(h, {"error": str(e)}, 500)

    return False

def handle_export(h, body):
    try:
        fmt = body.get("format", ""); raw_fp = body.get("path", ""); content = body.get("content", "")
        fp = safe_path(raw_fp) if raw_fp else ""
        out_dir = str(Path(fp).parent) if fp else HOME; stem = Path(fp).stem if fp else "export"

        if fmt == "text_to_pdf":
            from fpdf import FPDF
            pdf = FPDF(); pdf.add_page(); pdf.set_font("Helvetica", size=11)
            for line in (content or Path(fp).read_text(encoding="utf-8")).split("\n"):
                pdf.cell(0, 7, line[:120], ln=True)
            out = os.path.join(out_dir, stem + ".pdf"); pdf.output(out)
            return send_json(h, {"ok": True, "out": out})

        # Add other export formats here as needed (json_to_csv, etc.)
        return send_json(h, {"error": f"Noma'lum format: {fmt}"}, 400)
    except Exception as e: return send_json(h, {"error": str(e)}, 500)

def handle_excel_eval(h, body):
    try:
        formula = body.get("formula", "").strip()
        if not formula.startswith("="): return send_json(h, {"result": formula})
        import statistics as stats_mod
        vals = body.get("values", [])
        nums = [float(x) for x in vals if str(x).replace(".", "").replace("-", "").isdigit()]
        expr = formula[1:].upper(); result = "Xato"
        if expr.startswith("SUM"): result = sum(nums)
        elif expr.startswith("AVERAGE"): result = stats_mod.mean(nums) if nums else 0
        return send_json(h, {"result": result})
    except Exception as e: return send_json(h, {"error": str(e)}, 500)

def handle_docx_table(h, body):
    try:
        import docx
        fp = body.get("path", ""); rows = int(body.get("rows", 3)); cols = int(body.get("cols", 3))
        doc = docx.Document(fp); tbl = doc.add_table(rows=rows, cols=cols); tbl.style = "Table Grid"
        doc.save(fp); return send_json(h, {"ok": True})
    except Exception as e: return send_json(h, {"error": str(e)}, 500)

def handle_base64_op(h, body):
    try:
        action = body.get("action", "encode"); text = body.get("text", ""); res = ""
        if action == "encode": res = base64.b64encode(text.encode("utf-8")).decode()
        elif action == "decode": res = base64.b64decode(text).decode("utf-8", "replace")
        elif action == "encode_file": res = base64.b64encode(Path(body["path"]).read_bytes()).decode()
        return send_json(h, {"result": res})
    except Exception as e: return send_json(h, {"error": str(e)}, 500)
